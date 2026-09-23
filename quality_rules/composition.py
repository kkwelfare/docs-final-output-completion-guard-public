"""Hash-bound L0-L3 body-free composition evidence for docs quality receipts."""
from __future__ import annotations

import hashlib
import json
from itertools import combinations
from pathlib import Path
from typing import Any

import fitz

try:
    from . import semantic_surface
except ImportError:  # standalone plugin import path
    import semantic_surface  # type: ignore

SCHEMA_VERSION = "docs-composition-evidence-2"
SOURCE_GEOMETRY_SCHEMA = "docs-source-geometry-1"
PDF_TEXT_GEOMETRY_SCHEMA = "docs-pdf-text-geometry-1"
SPACING_PROJECTIONS_SCHEMA = "docs-spacing-projections-1"
SLIDE_KINDS = {"slide_deck", "slides", "training_slide_deck", "pptx_pdf"}
FOOTER_BAND_RATIO = 0.90
DEFAULT_MAX_REVIEW_GROUPS = 20
DEFAULT_MAX_SPACING_CANDIDATES = 5
SURFACE_STATE_SCHEMA = "docs-semantic-surface-policy-frames-1"
SURFACE_FIELDS = (
    "region_id", "page_index", "surface_tier", "topology", "container_scope",
    "visual_weight", "elevation_class", "semantic_group_id", "sequence_scope_id",
    "sequence_index",
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _handle(path: Path) -> dict[str, str]:
    return {"path": str(path.resolve()), "sha256": _sha256(path)}


def _entry(layer: str, rule_id: str, source_principle: str, measured: Any, threshold_or_preset: Any, status: str) -> dict[str, Any]:
    return {"layer": layer, "rule_id": rule_id, "source_principle": source_principle,
            "measured": measured, "threshold_or_preset": threshold_or_preset, "status": status}


def _num(value: Any) -> float | None:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _overlap(a: list[float], b: list[float]) -> float:
    x = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    y = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    return round(min(x, y) if x and y else 0.0, 3)


def _gutter(a: list[float], b: list[float]) -> float:
    horizontal = max(b[0] - a[2], a[0] - b[2], 0.0)
    vertical = max(b[1] - a[3], a[1] - b[3], 0.0)
    return round(max(horizontal, vertical), 3)


def _semantic_group(shape: Any) -> tuple[str | None, str | None]:
    """Read only explicit, non-body naming metadata: qa-semantic:<group>:<role>."""
    name = str(getattr(shape, "name", ""))
    prefix, group, role = "qa-semantic:", None, None
    if name.startswith(prefix):
        parts = name.split(":", 2)
        if len(parts) == 3 and parts[1].strip() and parts[2].strip():
            group, role = parts[1].strip(), parts[2].strip()
    return group, role


def _surface_region(shape: Any) -> str | None:
    """Read the producer's body-free shape marker: qa-surface:<region_id>:surface."""
    name = str(getattr(shape, "name", ""))
    parts = name.split(":")
    if len(parts) == 3 and parts[0] == "qa-surface" and parts[1].strip() and parts[2] == "surface":
        return parts[1].strip()
    return None


def _surface_state(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise ValueError("semantic surface state must be an object")
    result = {key: value[key] for key in SURFACE_FIELDS if key in value}
    if not isinstance(result.get("region_id"), str) or not isinstance(result.get("page_index"), int):
        raise ValueError("semantic surface state identity is malformed")
    return result


def _surface_context(producer_receipt: Any) -> dict[str, Any] | None:
    """Return validated body-free producer state, or None for every legacy caller."""
    if producer_receipt is None:
        return None
    if not isinstance(producer_receipt, dict):
        raise ValueError("semantic surface producer receipt is malformed")
    requested = producer_receipt.get("requested")
    declaration = requested.get("semantic_surface_policy") if isinstance(requested, dict) else None
    if declaration is None:
        return None
    if not isinstance(declaration, dict) or declaration.get("schema_version") != "docs-semantic-surface-policy-1":
        raise ValueError("semantic surface requested schema mismatch")
    identity = producer_receipt.get("execution_identity")
    if not isinstance(identity, dict) or set(identity) != {"contract_id", "policy_hash", "producer_run_id", "artifact_set_id"}:
        raise ValueError("semantic surface execution identity is malformed")
    states: dict[str, dict[str, dict[str, Any]]] = {}
    for state_name in ("applied", "effective"):
        aggregate = producer_receipt.get(state_name)
        if not isinstance(aggregate, dict) or aggregate.get("schema_version") != SURFACE_STATE_SCHEMA or aggregate.get("state") != state_name:
            raise ValueError(f"semantic surface {state_name} aggregate schema mismatch")
        frames = aggregate.get("frames")
        if not isinstance(frames, list):
            raise ValueError(f"semantic surface {state_name} frames are malformed")
        by_region: dict[str, dict[str, Any]] = {}
        for frame in frames:
            policy = frame.get("policy") if isinstance(frame, dict) else None
            state = _surface_state(policy)
            region_id = state["region_id"]
            if region_id in by_region:
                raise ValueError("semantic surface producer regions must be unique")
            by_region[region_id] = state
        states[state_name] = by_region
    requested_regions, tokens = declaration.get("regions"), declaration.get("tokens")
    if not isinstance(requested_regions, list) or not isinstance(tokens, list):
        raise ValueError("semantic surface requested regions are malformed")
    tokens_by_id = {token.get("id"): token for token in tokens if isinstance(token, dict) and isinstance(token.get("id"), str)}
    requested_by_region: dict[str, dict[str, Any]] = {}
    for region in requested_regions:
        if not isinstance(region, dict) or region.get("selected_token_id") not in tokens_by_id:
            raise ValueError("semantic surface requested token identity drift")
        state = _surface_state({**tokens_by_id[region["selected_token_id"]], **region})
        if state["region_id"] in requested_by_region:
            raise ValueError("semantic surface requested regions must be unique")
        requested_by_region[state["region_id"]] = state
    if set(requested_by_region) != set(states["applied"]) or set(requested_by_region) != set(states["effective"]):
        raise ValueError("semantic surface requested/applied/effective region identity drift")
    sequence_policy = declaration.get("sequence_policy")
    if sequence_policy is not None:
        if (not isinstance(sequence_policy, dict)
                or set(sequence_policy) != {"mode", "prominent_tiers", "max_consecutive_pages", "model_call_if_zero_candidates"}
                or sequence_policy.get("mode") != "review"
                or not isinstance(sequence_policy.get("prominent_tiers"), list)
                or not sequence_policy["prominent_tiers"]
                or any(not isinstance(item, str) or not item for item in sequence_policy["prominent_tiers"])
                or isinstance(sequence_policy.get("max_consecutive_pages"), bool)
                or not isinstance(sequence_policy.get("max_consecutive_pages"), int)
                or sequence_policy["max_consecutive_pages"] < 1
                or sequence_policy.get("model_call_if_zero_candidates") is not False):
            raise ValueError("semantic surface sequence policy is malformed")
        sequence_policy = {
            "mode": "review",
            "prominent_tiers": list(sequence_policy["prominent_tiers"]),
            "max_consecutive_pages": sequence_policy["max_consecutive_pages"],
            "model_call_if_zero_candidates": False,
        }
    effect_by_region: dict[str, dict[str, Any]] = {}
    effect_handle = producer_receipt.get("semantic_surface_effect_evidence")
    if effect_handle is not None:
        if (not isinstance(effect_handle, dict) or set(effect_handle) != {"path", "sha256", "size_bytes"}):
            raise ValueError("semantic surface effect evidence handle is malformed")
        effect_path = Path(str(effect_handle.get("path", "")))
        if (not effect_path.is_absolute() or not effect_path.is_file()
                or effect_handle.get("sha256") != _sha256(effect_path)
                or effect_handle.get("size_bytes") != effect_path.stat().st_size):
            raise ValueError("semantic surface effect evidence handle drift")
        try:
            effect_data = json.loads(effect_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            raise ValueError("semantic surface effect evidence is malformed") from exc
        if (not isinstance(effect_data, dict)
                or effect_data.get("schema_version") != "docs-semantic-surface-effect-evidence-1"
                or effect_data.get("execution_identity") != identity
                or not isinstance(effect_data.get("regions"), list)):
            raise ValueError("semantic surface effect evidence identity/schema drift")
        artifact_paths = {item.get("path"): item.get("sha256") for item in producer_receipt.get("artifacts", []) if isinstance(item, dict)}
        artifact_handle = effect_data.get("artifact")
        if not isinstance(artifact_handle, dict) or artifact_paths.get(artifact_handle.get("path")) != artifact_handle.get("sha256"):
            raise ValueError("semantic surface effect artifact hash drift")
        for item in effect_data["regions"]:
            region_id = item.get("region_id") if isinstance(item, dict) else None
            if not isinstance(region_id, str) or region_id in effect_by_region:
                raise ValueError("semantic surface effect region identity drift")
            effect_by_region[region_id] = item
    return {"execution_identity": identity, "requested": requested_by_region,
            "sequence_policy": sequence_policy,
            "metric_thresholds": declaration.get("metric_thresholds"),
            "effect_by_region": effect_by_region, **states}


def _role(shape: Any) -> str | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    text = getattr(getattr(shape, "text_frame", None), "text", "")
    if not isinstance(text, str) or not text.strip():
        return None
    if getattr(shape, "is_placeholder", False):
        kind = str(getattr(getattr(shape, "placeholder_format", None), "type", "")).lower()
        if "title" in kind:
            return "title"
    return "text-region"


def _line_boxes(artifact: Path) -> dict[int, list[list[float]]]:
    return {
        page: [item["bbox_pt"] for item in records]
        for page, records in _pdf_line_records(artifact).items()
    }


def _pdf_line_records(artifact: Path) -> dict[int, list[dict[str, Any]]]:
    """Extract body-free PDF line geometry with stable page/block/line ids."""
    result: dict[int, list[dict[str, Any]]] = {}
    with fitz.open(artifact) as pdf:
        for page_no, page in enumerate(pdf, start=1):
            lines: list[dict[str, Any]] = []
            for block_no, block in enumerate(page.get_text("dict").get("blocks", [])):
                if block.get("type") != 0:
                    continue
                for line_no, line in enumerate(block.get("lines", [])):
                    spans = [span for span in line.get("spans", []) if str(span.get("text", "")).strip()]
                    if spans:
                        bbox = [round(float(value), 3) for value in line["bbox"]]
                        origin = spans[0].get("origin")
                        baseline = origin[1] if isinstance(origin, (list, tuple)) and len(origin) >= 2 else bbox[3]
                        lines.append({
                            "line_id": f"p{page_no}-b{block_no}-l{line_no}",
                            "block_id": f"p{page_no}-b{block_no}",
                            "line_index": line_no,
                            "bbox_pt": bbox,
                            "baseline_y_pt": round(float(baseline), 3),
                            "line_height_pt": round(max(0.0, bbox[3] - bbox[1]), 3),
                            "span_count": len(spans),
                            "char_count": sum(len(str(span.get("text", ""))) for span in spans),
                            "font_size_pt": round(max(float(span.get("size", 0) or 0) for span in spans), 3),
                            "source_shape_id": None,
                        })
            result[page_no] = lines
    return result


def _vertical_occupancy(box: list[float], lines: list[list[float]]) -> float:
    spans = sorted((max(box[1], line[1]), min(box[3], line[3])) for line in lines
                   if line[2] > box[0] and line[0] < box[2] and line[3] > box[1] and line[1] < box[3])
    merged: list[list[float]] = []
    for start, end in spans:
        if end <= start:
            continue
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return round(sum(end - start for start, end in merged) / max(1.0, box[3] - box[1]), 5)


def _bottom_blank_ratio(box: list[float], lines: list[list[float]]) -> float:
    """Return empty vertical space after the final delivered line in a card."""
    bottoms = [min(box[3], line[3]) for line in lines
               if line[2] > box[0] and line[0] < box[2] and line[3] > box[1] and line[1] < box[3]]
    return round((box[3] - max(bottoms)) / max(1.0, box[3] - box[1]), 5) if bottoms else 1.0


def _box_area(box: list[float]) -> float:
    return max(0.0, box[2] - box[0]) * max(0.0, box[3] - box[1])


def _intersection_area(left: list[float], right: list[float]) -> float:
    width = max(0.0, min(left[2], right[2]) - max(left[0], right[0]))
    height = max(0.0, min(left[3], right[3]) - max(left[1], right[1]))
    return width * height


def _shape_kind(shape: Any) -> str:
    raw = getattr(getattr(shape, "shape_type", None), "name", None)
    value = str(raw or getattr(shape, "shape_type", "shape")).lower()
    if "picture" in value or "image" in value:
        return "image"
    if "table" in value:
        return "table"
    if "group" in value:
        return "group"
    if "connector" in value or "line" in value:
        return "line"
    if "text" in value:
        return "text_region"
    if "chart" in value:
        return "chart"
    return "shape"


def _line_source_region(line: dict[str, Any], regions: list[dict[str, Any]]) -> str | None:
    """Pair a delivered line with the smallest containing source text region."""
    bbox = line.get("bbox_pt")
    if not isinstance(bbox, list) or len(bbox) != 4:
        return None
    center = ((bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2)
    containing = [
        region for region in regions
        if region.get("shape_kind") == "text_region"
        and isinstance(region.get("bbox_pt"), list)
        and region["bbox_pt"][0] - 1 <= center[0] <= region["bbox_pt"][2] + 1
        and region["bbox_pt"][1] - 1 <= center[1] <= region["bbox_pt"][3] + 1
    ]
    if containing:
        return min(containing, key=lambda item: _box_area(item["bbox_pt"]))["shape_id"]
    overlaps = [
        region for region in regions
        if region.get("shape_kind") == "text_region"
        and isinstance(region.get("bbox_pt"), list)
        and _intersection_area(bbox, region["bbox_pt"]) > 0
    ]
    if overlaps:
        return max(
            overlaps,
            key=lambda item: _intersection_area(bbox, item["bbox_pt"]) / max(1.0, _box_area(bbox)),
        )["shape_id"]
    return None


def _union_boxes(boxes: list[list[float]]) -> list[float]:
    return [
        round(min(box[0] for box in boxes), 3),
        round(min(box[1] for box in boxes), 3),
        round(max(box[2] for box in boxes), 3),
        round(max(box[3] for box in boxes), 3),
    ]


def _projection_id(artifact: Path, kind: str, page: int, identity: Any) -> str:
    seed = json.dumps([_sha256(artifact), kind, page, identity], sort_keys=True, separators=(",", ":")).encode()
    return "spacing-" + hashlib.sha256(seed).hexdigest()[:16]


def _spacing_projection(
    artifact: Path, *, kind: str, rule: str, page: int, measurement: dict[str, Any],
    source_shape_ids: list[str], roles: dict[str, Any], relationship: dict[str, Any],
    bbox: list[float], pdf_line_ids: list[str] | None = None,
) -> dict[str, Any]:
    identity = [source_shape_ids, pdf_line_ids or [], measurement, roles, relationship, bbox]
    return {
        "id": _projection_id(artifact, kind, page, identity),
        "kind": kind,
        "rule": rule,
        "page": page,
        "units": "pt",
        "measurement_basis": "pdf_points",
        "measurement": {"units": "pt", **measurement},
        "source_shape_ids": list(source_shape_ids),
        "pdf_line_ids": list(pdf_line_ids or []),
        "roles": dict(roles),
        "relationship": dict(relationship),
        "bbox_pdf_points": list(bbox),
    }


def _build_spacing_projections(
    artifact: Path, source_pages: list[dict[str, Any]],
    pdf_lines: dict[int, list[dict[str, Any]]], page_sizes: dict[int, list[float]],
) -> dict[str, list[dict[str, Any]]]:
    """Build bounded, body-free measurements for Jev's qualitative reading."""
    line_baselines: list[dict[str, Any]] = []
    paragraph_gaps: list[dict[str, Any]] = []
    inner_margins: list[dict[str, Any]] = []
    outer_margins: list[dict[str, Any]] = []
    element_gaps: list[dict[str, Any]] = []
    source_by_page = {int(page["page"]): page.get("source_regions", []) for page in source_pages}
    for page_no, lines in pdf_lines.items():
        source_regions = [item for item in source_by_page.get(page_no, []) if isinstance(item, dict)]
        for line in lines:
            line["source_shape_id"] = _line_source_region(line, source_regions)
        ordered = sorted(lines, key=lambda item: (float(item["bbox_pt"][1]), float(item["bbox_pt"][0]), item["line_id"]))
        for previous, current in zip(ordered, ordered[1:]):
            previous_box, current_box = previous["bbox_pt"], current["bbox_pt"]
            bbox = _union_boxes([previous_box, current_box])
            distance = round(float(current["baseline_y_pt"]) - float(previous["baseline_y_pt"]), 3)
            box_gap = round(max(0.0, float(current_box[1]) - float(previous_box[3])), 3)
            source_ids = [item for item in (previous.get("source_shape_id"), current.get("source_shape_id")) if isinstance(item, str)]
            role_by_side = {
                "upper": next((r.get("semantic_role") or r.get("role") for r in source_regions if r.get("shape_id") == previous.get("source_shape_id")), "unknown"),
                "lower": next((r.get("semantic_role") or r.get("role") for r in source_regions if r.get("shape_id") == current.get("source_shape_id")), "unknown"),
            }
            line_baselines.append(_spacing_projection(
                artifact, kind="line_baseline_distance", rule="line_spacing_candidate", page=page_no,
                measurement={
                    "baseline_distance_pt": distance,
                    "line_box_gap_pt": box_gap,
                    "upper_line_height_pt": previous.get("line_height_pt"),
                    "lower_line_height_pt": current.get("line_height_pt"),
                },
                source_shape_ids=source_ids, pdf_line_ids=[previous["line_id"], current["line_id"]],
                roles=role_by_side,
                relationship={
                    "type": "same_block" if previous.get("block_id") == current.get("block_id") else "adjacent_pdf_blocks",
                    "block_id": previous.get("block_id"),
                    "upper_block_id": previous.get("block_id"),
                    "lower_block_id": current.get("block_id"),
                }, bbox=bbox,
            ))
            if previous.get("block_id") != current.get("block_id"):
                paragraph_gaps.append(_spacing_projection(
                    artifact, kind="paragraph_gap", rule="content_block_spacing_candidate", page=page_no,
                    measurement={
                        "paragraph_gap_pt": box_gap,
                        "baseline_distance_pt": distance,
                        "upper_line_height_pt": previous.get("line_height_pt"),
                        "lower_line_height_pt": current.get("line_height_pt"),
                    },
                    source_shape_ids=source_ids, pdf_line_ids=[previous["line_id"], current["line_id"]],
                    roles=role_by_side, relationship={"type": "adjacent_pdf_blocks", "upper_block_id": previous.get("block_id"), "lower_block_id": current.get("block_id")}, bbox=bbox,
                ))
        for region in source_regions:
            if region.get("shape_kind") != "text_region":
                continue
            matched = [line for line in lines if line.get("source_shape_id") == region.get("shape_id")]
            if not matched:
                continue
            line_box = _union_boxes([line["bbox_pt"] for line in matched])
            region_box = region["bbox_pt"]
            margins = {
                "left_pt": round(max(0.0, line_box[0] - region_box[0]), 3),
                "top_pt": round(max(0.0, line_box[1] - region_box[1]), 3),
                "right_pt": round(max(0.0, region_box[2] - line_box[2]), 3),
                "bottom_pt": round(max(0.0, region_box[3] - line_box[3]), 3),
            }
            inner_margins.append(_spacing_projection(
                artifact, kind="inner_margin", rule="content_block_intra_spacing_candidate", page=page_no,
                measurement={"inner_margins_pt": margins, "minimum_inner_margin_pt": min(margins.values()), "line_count": len(matched)},
                source_shape_ids=[region["shape_id"]], pdf_line_ids=[line["line_id"] for line in matched],
                roles={"region": region.get("semantic_role") or region.get("role") or "unknown"},
                relationship={"type": "source_region_to_pdf_lines", "source_region_id": region["shape_id"]},
                bbox=region_box,
            ))
        page_size = page_sizes.get(page_no)
        if page_size and ordered:
            line_box = _union_boxes([line["bbox_pt"] for line in ordered])
            margins = {
                "left_pt": round(max(0.0, line_box[0]), 3),
                "top_pt": round(max(0.0, line_box[1]), 3),
                "right_pt": round(max(0.0, page_size[0] - line_box[2]), 3),
                "bottom_pt": round(max(0.0, page_size[1] - line_box[3]), 3),
            }
            outer_margins.append(_spacing_projection(
                artifact, kind="outer_margin", rule="content_block_spacing_candidate", page=page_no,
                measurement={"outer_margins_pt": margins, "minimum_outer_margin_pt": min(margins.values()), "page_size_pt": list(page_size)},
                source_shape_ids=sorted({line["source_shape_id"] for line in ordered if isinstance(line.get("source_shape_id"), str)}),
                pdf_line_ids=[line["line_id"] for line in ordered],
                roles={"page": "page_boundary"}, relationship={"type": "pdf_lines_to_page_boundary"}, bbox=[0.0, 0.0, page_size[0], page_size[1]],
            ))
        # Only non-overlapping source regions participate in element gaps. Nested
        # regions are represented by inner margins, not as ambiguous gaps.
        for left_index, left in enumerate(source_regions):
            left_box = left.get("bbox_pt")
            if not isinstance(left_box, list):
                continue
            for right in source_regions[left_index + 1:]:
                right_box = right.get("bbox_pt")
                if not isinstance(right_box, list) or _intersection_area(left_box, right_box) > 0:
                    continue
                horizontal_overlap = max(0.0, min(left_box[2], right_box[2]) - max(left_box[0], right_box[0]))
                vertical_overlap = max(0.0, min(left_box[3], right_box[3]) - max(left_box[1], right_box[1]))
                if horizontal_overlap > 0 and left_box[3] <= right_box[1]:
                    gap, relation = right_box[1] - left_box[3], "below"
                elif horizontal_overlap > 0 and right_box[3] <= left_box[1]:
                    gap, relation = left_box[1] - right_box[3], "above"
                elif vertical_overlap > 0 and left_box[2] <= right_box[0]:
                    gap, relation = right_box[0] - left_box[2], "right_of"
                elif vertical_overlap > 0 and right_box[2] <= left_box[0]:
                    gap, relation = left_box[0] - right_box[2], "left_of"
                else:
                    continue
                element_gaps.append(_spacing_projection(
                    artifact, kind="element_gap", rule="content_block_spacing_candidate", page=page_no,
                    measurement={"element_gap_pt": round(max(0.0, gap), 3), "axis_overlap_pt": round(max(horizontal_overlap, vertical_overlap), 3)},
                    source_shape_ids=[left["shape_id"], right["shape_id"]],
                    roles={"first": left.get("semantic_role") or left.get("role") or "unknown", "second": right.get("semantic_role") or right.get("role") or "unknown"},
                    relationship={"type": relation, "first_shape_id": left["shape_id"], "second_shape_id": right["shape_id"]},
                    bbox=_union_boxes([left_box, right_box]),
                ))
    return {
        "line_baseline_distances": line_baselines,
        "paragraph_gaps": paragraph_gaps,
        "inner_margins": inner_margins,
        "outer_margins": outer_margins,
        "element_gaps": element_gaps,
    }


def _spacing_candidates(
    projections: dict[str, list[dict[str, Any]]], *, max_candidates: int,
) -> list[dict[str, Any]]:
    """Select one representative per measured spacing family before filling the cap."""
    ordered_families = (
        "line_baseline_distances", "paragraph_gaps", "inner_margins", "outer_margins", "element_gaps",
    )
    selected: list[dict[str, Any]] = []
    remaining: list[dict[str, Any]] = []
    for family in ordered_families:
        values = [item for item in projections.get(family, []) if isinstance(item, dict)]
        if values:
            selected.append(values[0])
            remaining.extend(values[1:])
    remaining.sort(key=lambda item: (int(item.get("page", 0)), str(item.get("kind", "")), str(item.get("id", ""))))
    selected.extend(remaining)
    candidates: list[dict[str, Any]] = []
    for projection in selected[:max(0, int(max_candidates))]:
        candidates.append({
            "id": projection["id"],
            "rule": projection["rule"],
            "severity": "review",
            "bucket": "REVIEW_CANDIDATE",
            "page": projection["page"],
            "shape_ids": list(projection.get("source_shape_ids", [])),
            "bbox_pdf_points": list(projection["bbox_pdf_points"]),
            "candidate_bbox_pdf_points": list(projection["bbox_pdf_points"]),
            "spacing_kind": projection["kind"],
            "spacing_projection_id": projection["id"],
            "spacing_projection": dict(projection),
            "source_member": (projection.get("source_shape_ids") or [None])[0],
            "paragraph": (projection.get("relationship") or {}).get("upper_block_id"),
            "candidate_type": projection["rule"],
            "units": "pt",
            "measurement_basis": projection.get("measurement_basis", "pdf_points"),
        })
    return candidates


def _contains(container: list[float], member: list[float]) -> bool:
    return member[0] >= container[0] and member[1] >= container[1] and member[2] <= container[2] and member[3] <= container[3]


def _semantic_card_rows(groups: dict[str, list[dict[str, Any]]], lines: list[list[float]], *,
                        min_occupancy: float, max_bottom_blank: float) -> list[dict[str, Any]]:
    """Evaluate only producer-opted-in cards; unlabeled sparse regions stay out of scope."""
    cards: list[dict[str, Any]] = []
    for group, members in sorted(groups.items()):
        by_role: dict[str, list[dict[str, Any]]] = {}
        for member in members:
            by_role.setdefault(str(member.get("semantic_role")), []).append(member)
        required = ("container", "heading", "body")
        malformed = [role for role in required if len(by_role.get(role, [])) != 1]
        containers = by_role.get("container", [])
        container = containers[0] if len(containers) == 1 else None
        out_of_container = [] if container is None else [member["shape_id"] for role in ("heading", "body")
            for member in by_role.get(role, []) if not _contains(container["bbox_pt"], member["bbox_pt"])]
        status = "block" if malformed or out_of_container else "pass"
        row: dict[str, Any] = {"group": group, "status": status, "malformed_roles": malformed,
            "out_of_container_shape_ids": out_of_container,
            "shape_ids": [member["shape_id"] for member in members]}
        if container is not None:
            box = container["bbox_pt"]
            occupancy = _vertical_occupancy(box, lines)
            bottom_blank = _bottom_blank_ratio(box, lines)
            row.update({"container_bbox_pt": box, "vertical_content_occupancy": occupancy,
                        "bottom_blank_ratio": bottom_blank, "min_card_vertical_occupancy": min_occupancy,
                        "max_card_bottom_blank_ratio": max_bottom_blank})
            if status == "pass" and occupancy < min_occupancy and bottom_blank > max_bottom_blank:
                row["status"] = "block"
                row["rule"] = "semantic_card_vertical_occupancy"
        cards.append(row)
    # Containers from different named cards must not overlap: nested/overlapping membership is ambiguous.
    for left, right in combinations(cards, 2):
        if "container_bbox_pt" in left and "container_bbox_pt" in right and _overlap(left["container_bbox_pt"], right["container_bbox_pt"]) > 0:
            left["status"] = right["status"] = "block"
            left["overlapping_group"] = right["group"]; right["overlapping_group"] = left["group"]
    return cards


def _candidate_id(artifact: Path, kind: str, page: int, shape_ids: list[str]) -> str:
    seed = json.dumps([_sha256(artifact), kind, page, shape_ids], sort_keys=True, separators=(",", ":")).encode()
    return "composition-" + hashlib.sha256(seed).hexdigest()[:16]


def _normalized_bbox(box: list[float], page_size: list[float]) -> list[float]:
    width, height = page_size
    return [round(box[0] / width, 2), round(box[1] / height, 2),
            round(box[2] / width, 2), round(box[3] / height, 2)]


def _band(value: float | None, step: float) -> float | None:
    return None if value is None else round(value / step) * step


def _review_group_key(raw: dict[str, Any], page_size: list[float]) -> tuple[Any, ...]:
    """Body-free stable grouping key; geometry and measurement bands only."""
    box = raw["candidate_bbox_pdf_points"]
    if raw["rule"] == "low_vertical_content_occupancy":
        return (raw["rule"], tuple(_normalized_bbox(box, page_size)),
                _band(_num(raw.get("vertical_content_occupancy")), 0.05))
    return (raw["rule"], tuple(_normalized_bbox(box, page_size)),
            _band(_num(raw.get("overlap_pt")), 2.0), _band(_num(raw.get("gutter_pt")), 2.0))


def _coalesce_reviews(artifact: Path, raw_candidates: list[dict[str, Any]], page_sizes: dict[int, list[float]], *, max_groups: int) -> list[dict[str, Any]]:
    grouped: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for raw in raw_candidates:
        page = raw.get("page")
        if not isinstance(page, int) or page not in page_sizes:
            continue
        grouped.setdefault(_review_group_key(raw, page_sizes[page]), []).append(raw)
    candidates: list[dict[str, Any]] = []
    for key in sorted(grouped, key=lambda value: json.dumps(value, separators=(",", ":"))):
        members = sorted(grouped[key], key=lambda item: (item["page"], item["shape_ids"], item["id"]))
        representative = members[0]
        affected_pages = sorted({member["page"] for member in members})
        shapes = sorted({shape for member in members for shape in member["shape_ids"]})
        candidate = dict(representative)
        candidate["id"] = _candidate_id(artifact, f"group:{representative['rule']}", representative["page"], representative["shape_ids"])
        candidate["affected_pages"] = affected_pages
        candidate["affected_count"] = len(members)
        candidate["example_shape_ids"] = shapes[:8]
        candidates.append(candidate)
    if len(candidates) <= max_groups:
        return candidates
    retained = candidates[:max_groups - 1]
    overflow = candidates[max_groups - 1:]
    pages = sorted({page for item in overflow for page in item["affected_pages"]})
    representative = overflow[0]
    retained.append({
        "id": _candidate_id(artifact, "review-group-overflow", representative["page"], representative["shape_ids"]),
        "rule": "composition_review_group_overflow",
        "page": representative["page"], "shape_ids": representative["shape_ids"],
        "candidate_bbox_pdf_points": representative["candidate_bbox_pdf_points"],
        "affected_pages": pages, "affected_count": sum(item["affected_count"] for item in overflow),
        "example_shape_ids": sorted({shape for item in overflow for shape in item["example_shape_ids"]})[:8],
        "total_group_count": len(candidates), "max_review_groups": max_groups,
    })
    return retained


def _surface_sequence_reviews(
    artifact: Path, surface_context: dict[str, Any],
    surface_shapes: dict[str, list[dict[str, Any]]], page_sizes: dict[int, list[float]],
) -> list[dict[str, Any]]:
    """Emit body-free review-only runs from the caller's bounded sequence policy."""
    policy = surface_context.get("sequence_policy")
    if not isinstance(policy, dict):
        return []
    prominent = set(policy["prominent_tiers"])
    limit = policy["max_consecutive_pages"]
    scoped: dict[str, dict[int, set[str]]] = {}
    for region_id, state in surface_context["effective"].items():
        scope_id, page_index = state.get("sequence_scope_id"), state.get("page_index")
        if (not isinstance(scope_id, str) or not scope_id or not isinstance(page_index, int)
                or state.get("surface_tier") not in prominent):
            continue
        scoped.setdefault(scope_id, {}).setdefault(page_index, set()).add(region_id)
    candidates: list[dict[str, Any]] = []
    for scope_id, region_ids_by_page in sorted(scoped.items()):
        pages = sorted(region_ids_by_page)
        runs: list[list[int]] = []
        for page_index in pages:
            if runs and page_index == runs[-1][-1] + 1:
                runs[-1].append(page_index)
            else:
                runs.append([page_index])
        for run in runs:
            if len(run) <= limit:
                continue
            affected_pages = [page_index + 1 for page_index in run]
            region_ids = sorted({region_id for page_index in run for region_id in region_ids_by_page[page_index]})
            shape_ids = sorted({shape["shape_id"] for region_id in region_ids
                                for shape in surface_shapes.get(region_id, [])})
            representative_page = affected_pages[0]
            page_size = page_sizes.get(representative_page)
            if page_size is None:
                continue
            seed = json.dumps(
                [_sha256(artifact), "SURFACE-SEQUENCE-REVIEW-01", scope_id, affected_pages, region_ids],
                sort_keys=True, separators=(",", ":"),
            ).encode()
            candidates.append({
                "id": "composition-" + hashlib.sha256(seed).hexdigest()[:16],
                "rule": "semantic_surface_prominent_tier_sequence",
                "rule_id": "SURFACE-SEQUENCE-REVIEW-01",
                "page": representative_page,
                "shape_ids": shape_ids,
                "candidate_bbox_pdf_points": [0.0, 0.0, page_size[0], page_size[1]],
                "sequence_scope_id": scope_id,
                "affected_pages": affected_pages,
                "affected_count": len(affected_pages),
                "example_shape_ids": shape_ids[:8],
                "region_ids": region_ids,
                "prominent_tiers": list(policy["prominent_tiers"]),
                "max_consecutive_pages": limit,
            })
    return candidates


def extract_evidence(source: Path, artifact: Path, *, minimum_gutter_pt: float | None = None,
                     max_review_groups: int = DEFAULT_MAX_REVIEW_GROUPS,
                     producer_receipt: dict[str, Any] | None = None,
                     include_spacing_candidates: bool = False,
                     max_spacing_candidates: int = DEFAULT_MAX_SPACING_CANDIDATES) -> dict[str, Any]:
    """Extract body-free PPTX/PDF evidence; no source or delivered body text is retained."""
    source, artifact = source.resolve(strict=True), artifact.resolve(strict=True)
    if source.suffix.lower() != ".pptx" or artifact.suffix.lower() != ".pdf":
        raise ValueError("composition extractor requires authoritative PPTX and delivered PDF")
    from pptx import Presentation

    presentation = Presentation(str(source))
    surface_context = _surface_context(producer_receipt)
    pdf_line_records = _pdf_line_records(artifact)
    pdf_lines = {page: [item["bbox_pt"] for item in records] for page, records in pdf_line_records.items()}
    pages: list[dict[str, Any]] = []
    hard_pairs: list[dict[str, Any]] = []
    raw_review_candidates: list[dict[str, Any]] = []
    page_sizes: dict[int, list[float]] = {}
    metrics: dict[str, float] = {}
    surface_shapes: dict[str, list[dict[str, Any]]] = {}
    surface_hard_failures: list[dict[str, Any]] = []
    source_sha256 = _sha256(source)
    width, height = float(presentation.slide_width) / 12700.0, float(presentation.slide_height) / 12700.0
    outer_margins: list[float] = []
    title_body_gaps: list[float] = []
    min_card_occupancy, max_card_bottom_blank = 0.35, 0.45
    for page_no, slide in enumerate(presentation.slides, start=1):
        regions: list[dict[str, Any]] = []
        source_regions: list[dict[str, Any]] = []
        semantic_groups: dict[str, list[dict[str, Any]]] = {}
        for shape_no, shape in enumerate(slide.shapes, start=1):
            group, declared_role = _semantic_group(shape)
            surface_region = _surface_region(shape)
            x, y = float(shape.left) / 12700.0, float(shape.top) / 12700.0
            w, h = float(shape.width) / 12700.0, float(shape.height) / 12700.0
            if w <= 0 or h <= 0:
                if surface_region is not None:
                    surface_hard_failures.append({"rule_id": "SURFACE-GEOMETRY-01", "region_id": surface_region,
                                                  "page_index": page_no - 1, "shape_id": f"p{page_no}-s{shape_no}"})
                continue
            box = [round(x, 3), round(y, 3), round(x + w, 3), round(y + h, 3)]
            shape_id = f"p{page_no}-s{shape_no}"
            shape_kind = _shape_kind(shape)
            text_role = _role(shape)
            source_regions.append({
                "shape_id": shape_id,
                "page": page_no,
                "page_index": page_no - 1,
                "shape_order": shape_no,
                "shape_kind": shape_kind,
                "role": declared_role or text_role or shape_kind,
                "semantic_role": declared_role,
                "semantic_group": group,
                "bbox_pt": box,
                "units": "pt",
                "geometry_basis": "pptx_points",
                "source_sha256": source_sha256,
                "has_text": bool(getattr(shape, "has_text_frame", False) and str(getattr(getattr(shape, "text_frame", None), "text", "")).strip()),
            })
            if surface_region is not None:
                surface_shapes.setdefault(surface_region, []).append({
                    "region_id": surface_region, "page_index": page_no - 1,
                    "shape_id": shape_id, "bbox_pt": box,
                })
                if box[0] < 0 or box[1] < 0 or box[2] > width or box[3] > height:
                    surface_hard_failures.append({"rule_id": "SURFACE-CONTAINMENT-01", "region_id": surface_region,
                                                  "page_index": page_no - 1, "shape_id": shape_id, "bbox_pt": box})
            if group is not None:
                semantic_groups.setdefault(group, []).append({"shape_id": shape_id, "semantic_role": declared_role, "bbox_pt": box})
            role = text_role
            if role is None:
                continue
            item = {"page": page_no, "shape_id": shape_id, "role": role, "bbox_pt": box}
            if group is not None:
                item["semantic_group"] = group
                item["semantic_role"] = declared_role
            regions.append(item)
            if h >= height * 0.55:
                occupancy = _vertical_occupancy(box, pdf_lines.get(page_no, []))
                if occupancy < 0.35:
                    raw_review_candidates.append({"id": _candidate_id(artifact, "low-vertical-occupancy", page_no, [item["shape_id"]]),
                        "rule": "low_vertical_content_occupancy", "page": page_no, "shape_ids": [item["shape_id"]],
                        "region_bbox_pt": box, "candidate_bbox_pdf_points": box, "vertical_content_occupancy": occupancy,
                        "minimum_occupancy": 0.35})
        semantic_cards = _semantic_card_rows(semantic_groups, pdf_lines.get(page_no, []),
                                              min_occupancy=min_card_occupancy, max_bottom_blank=max_card_bottom_blank)
        for left, right in combinations(regions, 2):
            overlap = _overlap(left["bbox_pt"], right["bbox_pt"])
            cy_left = (left["bbox_pt"][1] + left["bbox_pt"][3]) / 2
            cy_right = (right["bbox_pt"][1] + right["bbox_pt"][3]) / 2
            same_row = abs(cy_left - cy_right) <= min(left["bbox_pt"][3] - left["bbox_pt"][1], right["bbox_pt"][3] - right["bbox_pt"][1]) / 3
            gutter = _gutter(left["bbox_pt"], right["bbox_pt"]) if same_row else None
            explicit_siblings = left.get("semantic_group") and left.get("semantic_group") == right.get("semantic_group")
            record = {"page": page_no, "shape_ids": [left["shape_id"], right["shape_id"]], "overlap_pt": overlap,
                      "gutter_pt": gutter, "semantic_group": left.get("semantic_group") if explicit_siblings else None}
            if overlap > 0 or gutter is not None:
                if explicit_siblings:
                    hard_pairs.append(record)
                elif overlap > 0 or (minimum_gutter_pt is not None and gutter is not None and gutter < minimum_gutter_pt):
                    raw_review_candidates.append({"id": _candidate_id(artifact, "unclassified-overlap-or-gutter", page_no, record["shape_ids"]),
                        "rule": "unclassified_source_overlap_or_gutter", "page": page_no, "shape_ids": record["shape_ids"],
                        "candidate_bbox_pdf_points": [min(left["bbox_pt"][0], right["bbox_pt"][0]), min(left["bbox_pt"][1], right["bbox_pt"][1]), max(left["bbox_pt"][2], right["bbox_pt"][2]), max(left["bbox_pt"][3], right["bbox_pt"][3])],
                        "overlap_pt": overlap, "gutter_pt": gutter})
        titles = [region for region in regions if region["role"] == "title"]
        bodies = [region for region in regions if region["role"] == "text-region"]
        for title in titles:
            for body in bodies:
                if body["bbox_pt"][1] >= title["bbox_pt"][3]:
                    title_body_gaps.append(body["bbox_pt"][1] - title["bbox_pt"][3])
        # Prefer delivered non-footer PDF text lines; use non-full-bleed nonempty source regions only if PDF lines are absent.
        page_lines = [line for line in pdf_lines.get(page_no, []) if line[1] < height * FOOTER_BAND_RATIO]
        if page_lines:
            for line in page_lines:
                outer_margins.append(min(line[0], line[1], max(0.0, width - line[2]), max(0.0, height - line[3])))
        else:
            for region in regions:
                box = region["bbox_pt"]
                if box[1] < height * FOOTER_BAND_RATIO and box[0] > 0 and box[1] > 0 and box[2] < width and box[3] < height:
                    outer_margins.append(min(box[0], box[1], width - box[2], height - box[3]))
        page_size = [round(width, 3), round(height, 3)]
        page_sizes[page_no] = page_size
        pages.append({"page": page_no, "page_size_pt": page_size, "regions": regions,
                      "source_regions": source_regions, "semantic_cards": semantic_cards})
    if outer_margins:
        metrics["outer_margin_pt"] = round(min(outer_margins), 3)
    if title_body_gaps:
        metrics["title_body_gap_pt"] = round(min(title_body_gaps), 3)
    spacing_groups = _build_spacing_projections(artifact, pages, pdf_line_records, page_sizes)
    line_spacing = spacing_groups["line_baseline_distances"]
    paragraph_spacing = spacing_groups["paragraph_gaps"]
    inner_spacing = spacing_groups["inner_margins"]
    outer_spacing = spacing_groups["outer_margins"]
    element_spacing = spacing_groups["element_gaps"]
    if line_spacing:
        metrics["minimum_baseline_distance_pt"] = round(min(item["measurement"]["baseline_distance_pt"] for item in line_spacing), 3)
        metrics["minimum_line_box_gap_pt"] = round(min(item["measurement"]["line_box_gap_pt"] for item in line_spacing), 3)
    if paragraph_spacing:
        metrics["minimum_paragraph_gap_pt"] = round(min(item["measurement"]["paragraph_gap_pt"] for item in paragraph_spacing), 3)
    if inner_spacing:
        metrics["minimum_inner_margin_pt"] = round(min(item["measurement"]["minimum_inner_margin_pt"] for item in inner_spacing), 3)
    if outer_spacing:
        metrics["minimum_outer_margin_pt"] = round(min(item["measurement"]["minimum_outer_margin_pt"] for item in outer_spacing), 3)
    if element_spacing:
        metrics["minimum_element_gap_pt"] = round(min(item["measurement"]["element_gap_pt"] for item in element_spacing), 3)
    review_candidates = _coalesce_reviews(artifact, raw_review_candidates, page_sizes, max_groups=max(1, int(max_review_groups)))
    sequence_reviews = [] if surface_context is None else _surface_sequence_reviews(
        artifact, surface_context, surface_shapes, page_sizes,
    )
    review_candidates.extend(sequence_reviews)
    if include_spacing_candidates:
        existing_ids = {item.get("id") for item in review_candidates if isinstance(item, dict)}
        review_candidates.extend(
            item for item in _spacing_candidates(spacing_groups, max_candidates=max_spacing_candidates)
            if item.get("id") not in existing_ids
        )
    source_regions_flat = [
        region for page in pages for region in page.get("source_regions", [])
        if isinstance(region, dict)
    ]
    pdf_geometry = {
        "schema_version": PDF_TEXT_GEOMETRY_SCHEMA,
        "units": "pt",
        "measurement_basis": "pdf_points",
        "artifact": _handle(artifact),
        "pages": [
            {"page": page_no, "page_size_pt": page_sizes.get(page_no), "lines": records}
            for page_no, records in sorted(pdf_line_records.items())
        ],
    }
    source_geometry = {
        "schema_version": SOURCE_GEOMETRY_SCHEMA,
        "units": "pt",
        "measurement_basis": "pptx_points",
        "source": _handle(source),
        "pages": [
            {"page": page.get("page"), "page_size_pt": page.get("page_size_pt"), "regions": page.get("source_regions", [])}
            for page in pages
        ],
    }
    spacing_projection_data = {
        "schema_version": SPACING_PROJECTIONS_SCHEMA,
        "units": "pt",
        "measurement_basis": "pdf_points",
        "source": _handle(source),
        "artifact": _handle(artifact),
        **spacing_groups,
    }
    result = {"schema_version": SCHEMA_VERSION, "extractor": _handle(Path(__file__)), "source": _handle(source), "artifact": _handle(artifact),
            "units": "pt", "source_geometry": source_geometry, "source_regions": source_regions_flat,
            "pdf_text_geometry": pdf_geometry, "spacing_projections": spacing_projection_data,
            "pages": pages, "hard_pairs": hard_pairs, "review_candidates": review_candidates, "actual_metrics_pt": metrics,
            "limitations": ["Only non-whitespace PPTX text shapes participate in delivered line matching; body text is not retained.", "Source geometry retains shape roles and measured boxes, not document body or pixels.", "Only explicitly named qa-semantic:<group>:<role> siblings can hard-fail source geometry.", "Unclassified source relationships and unavailable metrics require bounded review, not a fabricated hard failure."]}
    if surface_context is not None:
        regions: list[dict[str, Any]] = []
        known = set(surface_context["requested"])
        for region_id in sorted(set(surface_shapes) - known):
            surface_hard_failures.append({"rule_id": "SURFACE-IDENTITY-01", "region_id": region_id,
                                          "observed_shape_ids": [item["shape_id"] for item in surface_shapes[region_id]]})
        for region_id in sorted(known):
            requested = surface_context["requested"][region_id]
            applied = surface_context["applied"][region_id]
            effective = surface_context["effective"][region_id]
            shapes = surface_shapes.get(region_id, [])
            expected_count = 0 if effective.get("topology") == "none" else 1
            if len(shapes) != expected_count:
                surface_hard_failures.append({"rule_id": "SURFACE-IDENTITY-01", "region_id": region_id,
                                              "expected_shape_count": expected_count, "observed_shape_count": len(shapes)})
            effect = surface_context.get("effect_by_region", {}).get(region_id)
            if surface_context.get("effect_by_region") and expected_count == 1 and effect is None:
                surface_hard_failures.append({"rule_id": "SURFACE-EFFECT-EVIDENCE-01", "region_id": region_id})
            for shape in shapes:
                if shape["page_index"] != effective["page_index"]:
                    surface_hard_failures.append({"rule_id": "SURFACE-CONTAINMENT-01", "region_id": region_id,
                                                  "expected_page_index": effective["page_index"],
                                                  "observed_page_index": shape["page_index"], "shape_id": shape["shape_id"]})
                regions.append({**shape, "requested": requested, "applied": applied, "effective": effective,
                                **({"effect_evidence": effect} if effect is not None else {})})
            if not shapes:
                regions.append({"region_id": region_id, "page_index": effective["page_index"], "shape_id": None,
                                "bbox_pt": None, "requested": requested, "applied": applied, "effective": effective})
        surface_ir = semantic_surface.build_ir(
            regions=regions, page_sizes=page_sizes, line_boxes=pdf_lines,
            hard_failures=surface_hard_failures,
            vertical_occupancy=_vertical_occupancy,
            bottom_blank_ratio=_bottom_blank_ratio,
            thresholds=(surface_context.get("metric_thresholds")
                        if isinstance(surface_context.get("metric_thresholds"), dict) else None),
        )
        existing_review_ids = {item.get("id") for item in result["review_candidates"] if isinstance(item, dict)}
        result["review_candidates"].extend(
            item for item in surface_ir["REVIEW_CANDIDATE"]
            if item.get("id") not in existing_review_ids
        )
        result["semantic_surfaces"] = {
            "schema_version": "docs-semantic-surface-composition-evidence-1",
            "execution_identity": surface_context["execution_identity"],
            "regions": regions,
            "HARD_FAIL": surface_hard_failures,
            "semantic_surface_ir": surface_ir,
            **({"sequence_review": {
                "mode": "review",
                "candidate_count": len(sequence_reviews),
                "model_call": False,
                "max_consecutive_pages": surface_context["sequence_policy"]["max_consecutive_pages"],
                "prominent_tiers": list(surface_context["sequence_policy"]["prominent_tiers"]),
            }} if isinstance(surface_context.get("sequence_policy"), dict) else {}),
        }
    return result


def write_evidence(source: Path, artifact: Path, evidence_dir: Path, *, minimum_gutter_pt: float | None = None,
                   max_review_groups: int = DEFAULT_MAX_REVIEW_GROUPS,
                   producer_receipt: dict[str, Any] | None = None,
                   include_spacing_candidates: bool = False,
                   max_spacing_candidates: int = DEFAULT_MAX_SPACING_CANDIDATES) -> Path:
    evidence_dir.mkdir(parents=True, exist_ok=True)
    path = evidence_dir / f"{artifact.name}.AQ-COMPOSITION-01.json"
    path.write_text(json.dumps(extract_evidence(source, artifact, minimum_gutter_pt=minimum_gutter_pt,
                                                 max_review_groups=max_review_groups,
                                                 producer_receipt=producer_receipt,
                                                 include_spacing_candidates=include_spacing_candidates,
                                                 max_spacing_candidates=max_spacing_candidates), ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")
    return path


def _load_evidence(value: Any, source: Path | None, artifact: Path | None) -> tuple[dict[str, Any] | None, str | None]:
    if not isinstance(value, dict): return None, "missing"
    raw, digest = value.get("path"), value.get("sha256")
    path = Path(raw) if isinstance(raw, str) else Path()
    if not path.is_absolute() or not path.is_file() or not isinstance(digest, str) or _sha256(path) != digest: return None, "missing-or-drift"
    try: data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError): return None, "invalid"
    if not isinstance(data, dict) or data.get("schema_version") != SCHEMA_VERSION: return None, "schema"
    if source is not None and data.get("source") != _handle(source): return None, "source-hash-drift"
    if artifact is not None and data.get("artifact") != _handle(artifact): return None, "artifact-hash-drift"
    return data, None


def evaluate(quality: Any, *, artifact_kind: str, baseline_statuses: Any = None, source: Path | None = None, artifact: Path | None = None) -> list[dict[str, Any]]:
    quality = quality if isinstance(quality, dict) else {}
    foundation, profile, preset = quality.get("quality_foundation"), quality.get("artifact_profile"), quality.get("design_preset")
    required = artifact_kind in SLIDE_KINDS or "quality_foundation" in quality
    if not required: return []
    rows: list[dict[str, Any]] = []
    baseline = baseline_statuses if isinstance(baseline_statuses, list) else []
    rows.append(_entry("L0", "L0-INTEGRITY-01", "existing entry rules after scan", baseline, "non-empty and all pass", "pass" if baseline and all(isinstance(x, dict) and x.get("status") == "pass" for x in baseline) else "block"))
    tokens = foundation.get("spacing_tokens_pt") if isinstance(foundation, dict) else None
    vals = list(tokens.values()) if isinstance(tokens, dict) else []
    rows.append(_entry("L1", "L1-SPACING-01", "spacing token consistency", len(vals), "non-empty positive tokens", "pass" if vals and all(_num(x) is not None and x > 0 for x in vals) else "block"))
    evidence, error = _load_evidence(quality.get("composition_evidence"), source, artifact)
    if evidence is None:
        rows.append(_entry("L1", "L1-REGION-01", "hash-bound semantic geometry", error, "required hash-bound evidence", "block"))
        rows.append(_entry("L2", "L2-MEASURE-01", "profile values compared with extractor metrics", error, "unavailable", "review"))
    else:
        semantic = foundation.get("semantic_regions") if isinstance(foundation, dict) else {}
        min_gutter = semantic.get("minimum_gutter_pt") if isinstance(semantic, dict) else None
        hard_pairs = [pair for pair in evidence.get("hard_pairs", []) if isinstance(pair, dict) and (_num(pair.get("overlap_pt")) and pair["overlap_pt"] > 0 or _num(min_gutter) is not None and _num(pair.get("gutter_pt")) is not None and pair["gutter_pt"] < min_gutter)]
        rows.append(_entry("L1", "L1-REGION-01", "explicit semantic sibling geometry", {"hard_pairs": hard_pairs, "evidence": quality.get("composition_evidence")}, {"minimum_gutter_pt": min_gutter}, "block" if hard_pairs else "pass"))
        surface = evidence.get("semantic_surfaces")
        if isinstance(surface, dict):
            surface_failures = surface.get("HARD_FAIL") if isinstance(surface.get("HARD_FAIL"), list) else [
                {"rule_id": "SURFACE-SCHEMA-01"}
            ]
            rows.append(_entry(
                "L1", "L1-SURFACE-01", "producer-bound semantic surface identity, containment, and geometry",
                {"region_count": len(surface.get("regions", [])) if isinstance(surface.get("regions"), list) else 0,
                 "HARD_FAIL": surface_failures, "evidence": quality.get("composition_evidence")},
                "no identity/hash/schema/containment/geometry violations",
                "block" if surface_failures else "pass",
            ))
        declared_values = profile.get("values") if isinstance(profile, dict) else {}
        card_min = _num(declared_values.get("min_card_vertical_occupancy")) if isinstance(declared_values, dict) else None
        card_max = _num(declared_values.get("max_card_bottom_blank_ratio")) if isinstance(declared_values, dict) else None
        if card_min is not None or card_max is not None:
            card_min = 0.35 if card_min is None else card_min
            card_max = 0.45 if card_max is None else card_max
            cards = [card for page in evidence.get("pages", []) if isinstance(page, dict)
                     for card in page.get("semantic_cards", []) if isinstance(card, dict)]
            violations = [card for card in cards if card.get("status") == "block" or
                          (_num(card.get("vertical_content_occupancy")) is not None and
                           _num(card.get("bottom_blank_ratio")) is not None and
                           card["vertical_content_occupancy"] < card_min and
                           card["bottom_blank_ratio"] > card_max)]
            rows.append(_entry("L1", "L1-SEMANTIC-CARD-01", "explicit semantic card occupancy from delivered PDF line boxes",
                               {"cards": cards, "violations": violations},
                               {"min_card_vertical_occupancy": card_min, "max_card_bottom_blank_ratio": card_max},
                               "block" if violations else "pass"))
        reviews = evidence.get("review_candidates", [])
        if reviews: rows.append(_entry("L1", "L1-COMPOSITION-REVIEW-01", "unclassified source relations and vertical content occupancy", reviews, "AQ-LAYOUT-01 hash-bound limited review", "review"))
        actual = evidence.get("actual_metrics_pt", {})
        declared = profile.get("values") if isinstance(profile, dict) else None
        if not isinstance(declared, dict):
            rows.append(_entry("L2", "L2-PROFILE-01", "artifact-class profile", "missing", "id and values required", "block" if artifact_kind in SLIDE_KINDS else "review"))
        else:
            rows.append(_entry("L2", "L2-PROFILE-01", "artifact-class profile", profile.get("id"), "id and values required", "pass" if isinstance(profile.get("id"), str) and profile["id"] else "block"))
            comparable = {key: actual[key] for key in declared if _num(actual.get(key)) is not None}
            failed = [key for key, value in comparable.items() if _num(declared.get(key)) is None or value < declared[key]]
            rows.append(_entry("L2", "L2-MEASURE-01", "profile thresholds against applicable extractor metrics", comparable if comparable else "unavailable", declared, "block" if failed else "pass" if comparable else "review"))
    if preset is None: rows.append(_entry("L3", "L3-PRESET-01", "master design choice", "absent", "optional", "pass"))
    else:
        needed, ident = isinstance(preset, dict) and preset.get("required") is True, preset.get("id") if isinstance(preset, dict) else None
        rows.append(_entry("L3", "L3-PRESET-01", "master design preset", ident, "id required only when required=true", "pass" if not needed or isinstance(ident, str) and ident else "block"))
    return rows


def blocking(rows: list[dict[str, Any]]) -> bool:
    return any(row.get("status") == "block" for row in rows)
