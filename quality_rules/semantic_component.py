"""Strict body-free semantic component policy, IR, and readback checks."""
from __future__ import annotations

import hashlib
import json
import math
from copy import deepcopy
from pathlib import Path
from typing import Any

POLICY_SCHEMA_VERSION = "docs-semantic-component-policy-1"
APPLIED_SCHEMA_VERSION = "docs-semantic-component-policy-applied-1"
EFFECTIVE_SCHEMA_VERSION = "docs-semantic-component-policy-effective-1"
IR_SCHEMA_VERSION = "docs-semantic-component-ir-1"
RECEIPT_SCHEMA_VERSION = "docs-semantic-component-producer-receipt-1"
BOUNDARY_POLICY_SCHEMA_VERSION = "docs-semantic-component-boundary-no-contact-policy-1"
BOUNDARY_EVIDENCE_SCHEMA_VERSION = "docs-semantic-component-boundary-no-contact-evidence-1"
BOUNDARY_REASON_CODE = "component_boundary_intersects_unowned_text"
FAMILIES = frozenset({"plain", "group", "card", "flow", "diagram", "table", "code", "callout"})
TOPOLOGIES = frozenset({"none", "stack", "paired", "sequence", "hierarchy", "grid", "band"})
DIVIDERS = frozenset({"none", "explicit-between", "grid-only"})
CONNECTORS = frozenset({"none", "directed", "undirected"})
ASSET_BACKENDS = frozenset({"none", "native", "svg"})
WRAPPER_BORDERS = frozenset({"none", "explicit"})
MEMBER_ROLES = frozenset({"container", "heading", "body", "item", "code", "asset"})
RELATION_KINDS = frozenset({"divider", "connector", "grid"})
IDENTITY_FIELDS = ("contract_id", "policy_hash", "producer_run_id", "artifact_set_id")
FORBIDDEN_BODY_KEYS = frozenset({"text", "body", "content", "payload", "raw", "value", "actual_value", "expected_value"})


def _canonical(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def digest(value: Any) -> str:
    return hashlib.sha256(_canonical(value)).hexdigest()


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _trimmed(value: Any, label: str) -> str:
    if not isinstance(value, str) or not value or value != value.strip():
        raise ValueError(f"{label} must be a trimmed non-empty string")
    return value


def _non_negative_int(value: Any, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ValueError(f"{label} must be a non-negative integer")
    return value


def _body_free(value: Any) -> bool:
    if isinstance(value, dict):
        return all(str(key).lower() not in FORBIDDEN_BODY_KEYS and _body_free(item) for key, item in value.items())
    if isinstance(value, list):
        return all(_body_free(item) for item in value)
    return True


def canonicalize_policy(value: Any) -> dict[str, Any]:
    required_fields = {"schema_version", "tokens", "components", "relations"}
    if not isinstance(value, dict) or not required_fields.issubset(value) or not set(value).issubset(required_fields | {"boundary_no_contact_policy"}):
        raise ValueError("semantic component policy declaration is malformed")
    if value.get("schema_version") != POLICY_SCHEMA_VERSION or not _body_free(value):
        raise ValueError("semantic component policy schema/body-free mismatch")
    raw_tokens = value.get("tokens")
    if not isinstance(raw_tokens, list) or not raw_tokens:
        raise ValueError("semantic component policy requires tokens")
    tokens: list[dict[str, Any]] = []
    token_by_id: dict[str, dict[str, Any]] = {}
    token_fields = {"id", "family", "topology", "divider_policy", "connector_policy", "asset_backend", "wrapper_border", "divider_placement", "divider_anchor"}
    for raw in raw_tokens:
        if not isinstance(raw, dict) or set(raw) != token_fields:
            raise ValueError("semantic component token is malformed")
        token_id = _trimmed(raw.get("id"), "semantic component token id")
        if token_id in token_by_id:
            raise ValueError("semantic component token IDs must be unique")
        token = {"id": token_id}
        for key, allowed in (("family", FAMILIES), ("topology", TOPOLOGIES), ("divider_policy", DIVIDERS),
                             ("connector_policy", CONNECTORS), ("asset_backend", ASSET_BACKENDS),
                             ("wrapper_border", WRAPPER_BORDERS)):
            if raw.get(key) not in allowed:
                raise ValueError(f"semantic component {key} token is unknown")
            token[key] = raw[key]
        if raw.get("divider_placement") != "after" or raw.get("divider_anchor") != "previous_block":
            raise ValueError("semantic component divider defaults must be after/previous_block")
        token["divider_placement"] = "after"; token["divider_anchor"] = "previous_block"
        if token["family"] == "plain" and any((token["divider_policy"] != "none", token["connector_policy"] != "none", token["wrapper_border"] != "none")):
            raise ValueError("plain component cannot imply relations or wrapper borders")
        if token["divider_policy"] == "grid-only" and (token["family"] != "table" or token["topology"] != "grid"):
            raise ValueError("grid-only divider requires table/grid")
        if token["wrapper_border"] == "explicit" and token["asset_backend"] != "svg":
            raise ValueError("explicit wrapper border requires svg backend")
        token_by_id[token_id] = token; tokens.append(token)

    raw_components = value.get("components")
    if not isinstance(raw_components, list) or not raw_components:
        raise ValueError("semantic component policy requires components")
    components: list[dict[str, Any]] = []
    component_by_id: dict[str, dict[str, Any]] = {}
    member_owner: dict[str, str] = {}
    component_fields = {"component_id", "caller_role", "page_index", "selected_token_id", "semantic_group_id", "members"}
    for raw in raw_components:
        if not isinstance(raw, dict) or set(raw) != component_fields:
            raise ValueError("semantic component declaration is malformed")
        component_id = _trimmed(raw.get("component_id"), "semantic component id")
        if component_id in component_by_id:
            raise ValueError("semantic component IDs must be unique")
        caller_role = _trimmed(raw.get("caller_role"), "semantic component caller role")
        page_index = _non_negative_int(raw.get("page_index"), "semantic component page_index")
        token_id = _trimmed(raw.get("selected_token_id"), "semantic component selected token")
        token = token_by_id.get(token_id)
        if token is None:
            raise ValueError("semantic component references unknown token")
        group_id = _trimmed(raw.get("semantic_group_id"), "semantic component group id")
        members_raw = raw.get("members")
        if not isinstance(members_raw, list) or not members_raw:
            raise ValueError("semantic component requires members")
        members=[]; roles=[]
        for member in members_raw:
            if not isinstance(member, dict) or set(member) != {"member_id", "role"}:
                raise ValueError("semantic component member is malformed")
            member_id = _trimmed(member.get("member_id"), "semantic member id")
            role = member.get("role")
            if role not in MEMBER_ROLES or member_id in member_owner:
                raise ValueError("semantic member role/identity is invalid")
            member_owner[member_id] = component_id; roles.append(role); members.append({"member_id": member_id, "role": role})
        if token["family"] == "card" and not {"container", "heading", "body"}.issubset(roles):
            raise ValueError("semantic card requires container heading body members")
        component = {"component_id": component_id, "caller_role": caller_role, "page_index": page_index,
                     "selected_token_id": token_id, "semantic_group_id": group_id,
                     "members": sorted(members, key=lambda item: item["member_id"])}
        component_by_id[component_id] = component; components.append(component)

    relations: list[dict[str, Any]] = []
    seen_relations: set[str] = set(); seen_boundaries: set[tuple[str, str, str]] = set()
    for raw in value.get("relations", []):
        if not isinstance(raw, dict):
            raise ValueError("semantic relation is malformed")
        required = {"relation_id", "kind", "from_member_id", "to_member_id"}
        optional = {"placement", "anchor", "override_reason", "boundary_id"}
        if not required.issubset(raw) or not set(raw).issubset(required | optional):
            raise ValueError("semantic relation fields are malformed")
        relation_id = _trimmed(raw.get("relation_id"), "semantic relation id")
        if relation_id in seen_relations:
            raise ValueError("semantic relation IDs must be unique")
        kind = raw.get("kind")
        if kind not in RELATION_KINDS:
            raise ValueError("semantic relation kind is unknown")
        left = _trimmed(raw.get("from_member_id"), "semantic relation from member")
        right = _trimmed(raw.get("to_member_id"), "semantic relation to member")
        if left == right or left not in member_owner or right not in member_owner:
            raise ValueError("semantic relation members are missing")
        owners = {component_by_id[member_owner[left]]["selected_token_id"], component_by_id[member_owner[right]]["selected_token_id"]}
        if kind == "divider" and any(token_by_id[token]["divider_policy"] != "explicit-between" for token in owners):
            raise ValueError("undeclared divider relation")
        if kind == "connector" and any(token_by_id[token]["connector_policy"] == "none" for token in owners):
            raise ValueError("undeclared connector relation")
        if kind == "grid" and any(token_by_id[token]["divider_policy"] != "grid-only" for token in owners):
            raise ValueError("undeclared grid relation")
        placement = raw.get("placement", "after"); anchor = raw.get("anchor", "previous_block")
        reason = raw.get("override_reason")
        if placement == "before":
            if not isinstance(reason, str) or not reason.strip():
                raise ValueError("before divider requires explicit override reason")
        elif placement != "after" or anchor != "previous_block" or reason is not None:
            raise ValueError("semantic divider placement/anchor is invalid")
        boundary_id = _trimmed(raw.get("boundary_id", relation_id), "semantic relation boundary id")
        boundary = (kind, boundary_id, "|".join(sorted((left, right))))
        if boundary in seen_boundaries:
            raise ValueError("duplicate semantic boundary")
        seen_boundaries.add(boundary); seen_relations.add(relation_id)
        item={"relation_id":relation_id,"kind":kind,"from_member_id":left,"to_member_id":right,
              "placement":placement,"anchor":anchor,"boundary_id":boundary_id}
        if reason is not None: item["override_reason"] = reason.strip()
        relations.append(item)
    result={"schema_version": POLICY_SCHEMA_VERSION,
            "tokens": sorted(tokens, key=lambda item: item["id"]),
            "components": sorted(components, key=lambda item: item["component_id"]),
            "relations": sorted(relations, key=lambda item: item["relation_id"])}
    boundary_policy=value.get("boundary_no_contact_policy")
    if boundary_policy is not None:
        required_boundary_fields={"schema_version","status","surface_families","allowed_overlay_ids"}
        if (not isinstance(boundary_policy,dict) or not required_boundary_fields.issubset(boundary_policy)
                or not set(boundary_policy).issubset(required_boundary_fields | {"member_ownership"})
                or boundary_policy.get("schema_version")!=BOUNDARY_POLICY_SCHEMA_VERSION
                or boundary_policy.get("status")!="declared"):
            raise ValueError("semantic component boundary no-contact policy is malformed")
        families=boundary_policy.get("surface_families")
        overlays=boundary_policy.get("allowed_overlay_ids")
        if (not isinstance(families,list) or not families or len(set(families))!=len(families)
                or any(item not in {"card","group"} for item in families)):
            raise ValueError("semantic component boundary surface families are malformed")
        if (not isinstance(overlays,list) or len(set(overlays))!=len(overlays)
                or any(not isinstance(item,str) or not item or item!=item.strip() for item in overlays)):
            raise ValueError("semantic component boundary overlay IDs are malformed")
        ownership=boundary_policy.get("member_ownership",[]); normalized_ownership=[]; ownership_ids=set()
        if not isinstance(ownership,list): raise ValueError("semantic component boundary member ownership is malformed")
        for item in ownership:
            if not isinstance(item,dict) or set(item)!={"text_block_id","component_id"}:
                raise ValueError("semantic component boundary member ownership is malformed")
            text_block_id=_trimmed(item.get("text_block_id"),"semantic component boundary ownership text block")
            component_id=_trimmed(item.get("component_id"),"semantic component boundary ownership component")
            if text_block_id in ownership_ids or component_id not in component_by_id:
                raise ValueError("semantic component boundary member ownership is malformed")
            ownership_ids.add(text_block_id); normalized_ownership.append({"text_block_id":text_block_id,"component_id":component_id})
        result["boundary_no_contact_policy"]={"schema_version":BOUNDARY_POLICY_SCHEMA_VERSION,"status":"declared",
            "surface_families":sorted(families),"allowed_overlay_ids":sorted(overlays)}
        if "member_ownership" in boundary_policy:
            result["boundary_no_contact_policy"]["member_ownership"]=sorted(normalized_ownership,key=lambda item:item["text_block_id"])
    return result


def resolve_component(policy: Any, *, component_id: str, caller_role: str,
                      direct_values: dict[str, Any] | None = None) -> tuple[dict[str, Any], dict[str, Any]]:
    declaration = canonicalize_policy(policy)
    component_id = _trimmed(component_id, "semantic component id")
    caller_role = _trimmed(caller_role, "semantic component caller role")
    component = {item["component_id"]: item for item in declaration["components"]}.get(component_id)
    if component is None or component["caller_role"] != caller_role:
        raise ValueError("semantic component identity/role drift")
    token = {item["id"]: item for item in declaration["tokens"]}[component["selected_token_id"]]
    requested={**deepcopy(component), **{key: token[key] for key in token if key != "id"}}
    supplied = {} if direct_values is None else direct_values
    if not isinstance(supplied, dict) or any(key not in requested for key in supplied):
        raise ValueError("semantic component direct values are malformed")
    for key, item in supplied.items():
        if item != requested[key]:
            raise ValueError(f"semantic component direct-value drift: {key}")
    applied={"schema_version":APPLIED_SCHEMA_VERSION, **deepcopy(requested)}
    return requested, applied


def execution_identity(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict) or set(value) != set(IDENTITY_FIELDS):
        raise ValueError("semantic component execution identity must contain four fields")
    for key in ("contract_id", "policy_hash", "artifact_set_id"):
        token=value.get(key)
        if not isinstance(token,str) or len(token)!=64 or any(ch not in "0123456789abcdef" for ch in token):
            raise ValueError("semantic component execution identity hash is invalid")
    _non_negative_int(value.get("producer_run_id"), "semantic component producer_run_id")
    return {key:value[key] for key in IDENTITY_FIELDS}


def build_frame(*, policy: dict[str, Any], component_id: str, caller_role: str,
                identity: dict[str, Any], format_name: str, applied_geometry: dict[str, Any],
                effective_geometry: dict[str, Any], direct_values: dict[str, Any] | None = None) -> dict[str, Any]:
    requested, applied = resolve_component(policy, component_id=component_id, caller_role=caller_role, direct_values=direct_values)
    if format_name not in {"pptx", "svg", "html"} or not _body_free(applied_geometry) or not _body_free(effective_geometry):
        raise ValueError("semantic component geometry/format is invalid")
    effective={"schema_version":EFFECTIVE_SCHEMA_VERSION, **deepcopy(requested), "format":format_name,
               "geometry":deepcopy(effective_geometry)}
    frame={"schema_version":"docs-semantic-component-frame-1","execution_identity":execution_identity(identity),
           "requested":requested,"applied":{**applied,"format":format_name,"geometry":deepcopy(applied_geometry)},
           "effective":effective}
    validate_frame(frame, policy)
    return frame


def _finite(value: Any) -> bool:
    return not isinstance(value,bool) and isinstance(value,(int,float)) and math.isfinite(float(value))


def validate_frame(frame: Any, policy: dict[str, Any]) -> None:
    if not isinstance(frame,dict) or set(frame)!={"schema_version","execution_identity","requested","applied","effective"} or not _body_free(frame):
        raise ValueError("semantic component frame schema/body-free mismatch")
    execution_identity(frame["execution_identity"])
    requested=frame["requested"]; applied=frame["applied"]; effective=frame["effective"]
    expected, expected_applied=resolve_component(policy,component_id=requested.get("component_id"),caller_role=requested.get("caller_role"))
    if requested!=expected or any(applied.get(k)!=v for k,v in expected_applied.items()) or any(effective.get(k)!=v for k,v in expected.items()):
        raise ValueError("semantic component requested/applied/effective identity drift")
    geometry=effective.get("geometry")
    if not isinstance(geometry,dict): raise ValueError("semantic component effective geometry missing")
    if expected["family"]=="card":
        names=geometry.get("shape_names")
        required={f"qa-semantic:{expected['semantic_group_id']}:{role}" for role in ("container","heading","body")}
        if not isinstance(names,list) or not required.issubset(names): raise ValueError("semantic card member naming incomplete")
        if geometry.get("surface_present") is not True: raise ValueError("semantic component surface mismatch")
    if expected["divider_policy"]=="none" and geometry.get("divider_count",0)!=0:
        raise ValueError("undeclared divider emitted")
    if expected["wrapper_border"]=="none" and geometry.get("wrapper_border_count",0)!=0:
        raise ValueError("unwanted SVG wrapper border")
    for divider in geometry.get("dividers",[]):
        if (not isinstance(divider,dict) or divider.get("foreground") is True or divider.get("text_bbox_crossing") is True
                or not _finite(divider.get("block_to_divider_gap_pt")) or not _finite(divider.get("divider_to_next_block_gap_pt"))
                or float(divider["block_to_divider_gap_pt"]) >= float(divider["divider_to_next_block_gap_pt"])):
            raise ValueError("semantic divider rendered geometry mismatch")
    for connector in geometry.get("connectors",[]):
        if (not isinstance(connector,dict) or connector.get("behind_surfaces") is not True
                or connector.get("endpoint_clipped_to_boundary") is not True
                or connector.get("visible_text_crossing_count",0)!=0 or connector.get("visible_card_interior_length_pt",0)!=0):
            raise ValueError("semantic connector rendered geometry mismatch")


def _bbox(value: Any, label: str) -> list[float]:
    if (not isinstance(value,list) or len(value)!=4 or any(not _finite(item) for item in value)
            or float(value[2]) < float(value[0]) or float(value[3]) < float(value[1])):
        raise ValueError(f"semantic component {label} bbox is malformed")
    return [round(float(item),3) for item in value]


def _handle(path: Path) -> dict[str, Any]:
    resolved=path.resolve(strict=True)
    return {"path":str(resolved),"sha256":sha256(resolved),"size_bytes":resolved.stat().st_size}


def _perimeter_intersections(surface: list[float], text: list[float]) -> list[dict[str,Any]]:
    sx0,sy0,sx1,sy1=surface; tx0,ty0,tx1,ty1=text
    found=[]
    for edge,y in (("top",sy0),("bottom",sy1)):
        start=max(sx0,tx0); end=min(sx1,tx1)
        if ty0 <= y <= ty1 and end > start:
            found.append({"edge":edge,"start_pt":[round(start,3),round(y,3)],"end_pt":[round(end,3),round(y,3)],"length_pt":round(end-start,3)})
    for edge,x in (("left",sx0),("right",sx1)):
        start=max(sy0,ty0); end=min(sy1,ty1)
        if tx0 <= x <= tx1 and end > start:
            found.append({"edge":edge,"start_pt":[round(x,3),round(start,3)],"end_pt":[round(x,3),round(end,3)],"length_pt":round(end-start,3)})
    return found


def evaluate_boundary_no_contact(*, policy: dict[str,Any], identity: dict[str,Any], page_index: int,
                                 surfaces: list[dict[str,Any]], text_blocks: list[dict[str,Any]],
                                 source_artifact: Path, rendered_artifact: Path) -> dict[str,Any]:
    """Evaluate the format-neutral perimeter↔unowned-text relation with rendered coverage."""
    declaration=canonicalize_policy(policy); identity_value=execution_identity(identity)
    page_index=_non_negative_int(page_index,"semantic component boundary page_index")
    source_handle=_handle(source_artifact); rendered_handle=_handle(rendered_artifact)
    boundary_policy=declaration.get("boundary_no_contact_policy")
    base={"schema_version":BOUNDARY_EVIDENCE_SCHEMA_VERSION,"execution_identity":identity_value,
          "policy_state":{"requested":"enforced","applied":"enforced","effective":"enforced"},
          "source_artifact":source_handle,"rendered_artifact":rendered_handle,"page_index":page_index}
    if boundary_policy is None:
        return {**base,"policy_state":{"requested":"undeclared","applied":"not_applicable","effective":"not_applicable"},
                "status":"not_applicable","reason_code":None,"violation_count":0,"duplicate_candidate_count":0,"violations":[]}
    if not isinstance(surfaces,list) or not isinstance(text_blocks,list):
        raise ValueError("semantic component boundary primitives are malformed")
    components={item["component_id"]:item for item in declaration["components"]}
    tokens={item["id"]:item for item in declaration["tokens"]}
    eligible={component_id for component_id,item in components.items()
              if item["page_index"]==page_index and tokens[item["selected_token_id"]]["family"] in boundary_policy["surface_families"]}
    normalized_surfaces=[]
    for raw in surfaces:
        fields={"component_id","surface_member_id","source_primitive_id","bbox_pt","outlined"}
        if not isinstance(raw,dict) or set(raw)!=fields or not isinstance(raw.get("outlined"),bool):
            raise ValueError("semantic component boundary surface primitive is malformed")
        component_id=_trimmed(raw.get("component_id"),"semantic component boundary surface owner")
        if component_id not in eligible:
            raise ValueError("semantic component boundary surface ownership/page drift")
        member_id=_trimmed(raw.get("surface_member_id"),"semantic component boundary surface member")
        component=components[component_id]
        family=tokens[component["selected_token_id"]]["family"]
        member_matches=any(item["member_id"]==member_id and item["role"]=="container" for item in component["members"])
        group_surface_matches=family=="group" and member_id==f"{component_id}-container"
        if not (member_matches or group_surface_matches):
            raise ValueError("semantic component boundary surface ownership is malformed")
        normalized_surfaces.append({"component_id":component_id,"surface_member_id":member_id,
            "source_primitive_id":_trimmed(raw.get("source_primitive_id"),"semantic component boundary source primitive"),
            "bbox_pt":_bbox(raw.get("bbox_pt"),"surface"),"outlined":raw["outlined"]})
    overlays=set(boundary_policy["allowed_overlay_ids"]); normalized_text=[]
    for raw in text_blocks:
        fields={"text_block_id","owner_component_id","source_primitive_id","bbox_pt","rendered_coverage_bbox_pt","explicit_overlay"}
        if not isinstance(raw,dict) or set(raw)!=fields or not isinstance(raw.get("explicit_overlay"),bool):
            raise ValueError("semantic component text block primitive is malformed")
        block_id=_trimmed(raw.get("text_block_id"),"semantic component text block id")
        owner=raw.get("owner_component_id")
        if owner is not None and owner not in components:
            raise ValueError("semantic component text ownership is malformed")
        if raw["explicit_overlay"] and block_id not in overlays:
            raise ValueError("semantic component overlay permission is malformed")
        normalized_text.append({"text_block_id":block_id,"owner_component_id":owner,
            "source_primitive_id":_trimmed(raw.get("source_primitive_id"),"semantic component text source primitive"),
            "bbox_pt":_bbox(raw.get("bbox_pt"),"text block"),
            "rendered_coverage_bbox_pt":_bbox(raw.get("rendered_coverage_bbox_pt"),"rendered coverage"),
            "explicit_overlay":raw["explicit_overlay"]})
    violations=[]; candidate_ids=set(); duplicates=0
    for surface in normalized_surfaces:
        if not surface["outlined"]: continue
        for text in normalized_text:
            if text["owner_component_id"] is not None or text["explicit_overlay"]: continue
            intersections=_perimeter_intersections(surface["bbox_pt"],text["bbox_pt"])
            if not intersections: continue
            intersection=max(intersections,key=lambda item:(item["length_pt"],item["edge"]))
            candidate_id=digest({"page_index":page_index,"component_id":surface["component_id"],
                "text_block_id":text["text_block_id"],"edge":intersection["edge"],"start_pt":intersection["start_pt"],"end_pt":intersection["end_pt"]})
            if candidate_id in candidate_ids:
                duplicates+=1; continue
            candidate_ids.add(candidate_id)
            violations.append({"candidate_id":candidate_id,"reason_code":BOUNDARY_REASON_CODE,"page_index":page_index,
                "component_id":surface["component_id"],"surface_member_id":surface["surface_member_id"],
                "surface_source_primitive_id":surface["source_primitive_id"],"surface_bbox_pt":surface["bbox_pt"],
                "text_block_id":text["text_block_id"],"text_source_primitive_id":text["source_primitive_id"],
                "text_bbox_pt":text["bbox_pt"],"rendered_coverage_bbox_pt":text["rendered_coverage_bbox_pt"],
                "intersection":intersection})
    violations.sort(key=lambda item:(item["page_index"],item["component_id"],item["text_block_id"],item["intersection"]["edge"]))
    return {**base,"status":"fail" if violations else "pass","reason_code":BOUNDARY_REASON_CODE if violations else None,
            "violation_count":len(violations),"duplicate_candidate_count":duplicates,"violations":violations}


def analyze_pptx_pdf_boundary_contacts(*, execution_contract: Path | dict[str,Any], pptx_path: Path,
                                       pdf_path: Path, page_indexes: list[int] | None=None) -> dict[str,Any]:
    """Adapt actual PPTX/PDF geometry after binding policy and identity to one live contract."""
    from quality_rules import execution_contract as execution_contract_rules
    from pptx import Presentation
    import fitz
    contract, errors=execution_contract_rules.load_contract(execution_contract,require_artifacts=True)
    if errors or contract is None:
        raise ValueError("semantic component boundary execution contract is invalid: " + "; ".join(errors))
    full_policy=contract["declaration"]["policy"]
    policy=full_policy.get("semantic_component_policy")
    if not isinstance(policy,dict):
        raise ValueError("semantic component boundary execution contract does not declare semantic component policy")
    declaration=canonicalize_policy(policy)
    if policy!=declaration:
        raise ValueError("semantic component boundary execution contract semantic policy is not canonical")
    identity_value=execution_contract_rules.identity(contract)
    pptx=pptx_path.resolve(strict=True); pdf=pdf_path.resolve(strict=True)
    artifact_paths={item["path"] for item in contract["artifacts"]}
    if str(pptx) not in artifact_paths or str(pdf) not in artifact_paths:
        raise ValueError("semantic component boundary PPTX/PDF artifacts are not bound by execution contract")
    presentation=Presentation(pptx); document=fitz.open(pdf)
    if len(presentation.slides)!=len(document):
        raise ValueError("semantic component PPTX/PDF page identity drift")
    selected=list(range(len(document))) if page_indexes is None else page_indexes
    if (not isinstance(selected,list) or len(set(selected))!=len(selected)
            or any(isinstance(item,bool) or not isinstance(item,int) or item<0 or item>=len(document) for item in selected)):
        raise ValueError("semantic component boundary page selection is malformed")
    boundary_policy=declaration.get("boundary_no_contact_policy")
    if boundary_policy is None:
        return evaluate_boundary_no_contact(policy=declaration,identity=identity_value,page_index=selected[0] if selected else 0,
            surfaces=[],text_blocks=[],source_artifact=pptx,rendered_artifact=pdf)
    tokens={item["id"]:item for item in declaration["tokens"]}; page_results=[]
    for page_index in selected:
        slide=presentation.slides[page_index]; page=document[page_index]
        scale_x=page.rect.width/(presentation.slide_width/12700); scale_y=page.rect.height/(presentation.slide_height/12700)
        components=[item for item in declaration["components"] if item["page_index"]==page_index
                    and tokens[item["selected_token_id"]]["family"] in boundary_policy["surface_families"]]
        owner_by_name={f"qa-semantic:{item['semantic_group_id']}:{member['role']}":item["component_id"]
                       for item in declaration["components"] for member in item["members"]}
        owner_by_name.update({item["text_block_id"]:item["component_id"] for item in boundary_policy.get("member_ownership",[])})
        shapes_by_name={shape.name:shape for shape in slide.shapes}; surfaces=[]
        for component in components:
            name=f"qa-semantic:{component['semantic_group_id']}:container"; shape=shapes_by_name.get(name)
            if shape is None: continue
            container=next((item for item in component["members"] if item["role"]=="container"),None)
            member_id=container["member_id"] if container is not None else f"{component['component_id']}-container"
            fill_type=getattr(getattr(shape.line,"fill",None),"type",None)
            surfaces.append({"component_id":component["component_id"],"surface_member_id":member_id,
                "source_primitive_id":name,"bbox_pt":[round(value/12700,3) for value in (shape.left,shape.top,shape.left+shape.width,shape.top+shape.height)],
                "outlined":getattr(fill_type,"name",str(fill_type))=="SOLID"})
        words=page.get_text("words"); text_blocks=[]; overlays=set(boundary_policy["allowed_overlay_ids"])
        for shape in slide.shapes:
            if not getattr(shape,"has_text_frame",False) or not shape.text.strip(): continue
            bbox=[shape.left/12700,shape.top/12700,(shape.left+shape.width)/12700,(shape.top+shape.height)/12700]
            pdf_bbox=[bbox[0]*scale_x,bbox[1]*scale_y,bbox[2]*scale_x,bbox[3]*scale_y]
            covered=[word for word in words if pdf_bbox[0] <= (word[0]+word[2])/2 <= pdf_bbox[2]
                     and pdf_bbox[1] <= (word[1]+word[3])/2 <= pdf_bbox[3]]
            if not covered:
                potential=any(_perimeter_intersections([round(v/12700,3) for v in (surface.left,surface.top,surface.left+surface.width,surface.top+surface.height)],bbox)
                              for surface in slide.shapes if surface.name.startswith("qa-semantic:") and surface.name.endswith(":container"))
                if potential: raise ValueError("semantic component rendered coverage missing for intersecting text block")
                continue
            coverage=[min(item[0] for item in covered),min(item[1] for item in covered),max(item[2] for item in covered),max(item[3] for item in covered)]
            text_blocks.append({"text_block_id":shape.name,"owner_component_id":owner_by_name.get(shape.name),
                "source_primitive_id":shape.name,"bbox_pt":[round(item,3) for item in bbox],
                "rendered_coverage_bbox_pt":[round(item,3) for item in coverage],"explicit_overlay":shape.name in overlays})
        page_results.append(evaluate_boundary_no_contact(policy=declaration,identity=identity_value,page_index=page_index,
            surfaces=surfaces,text_blocks=text_blocks,source_artifact=pptx,rendered_artifact=pdf))
    violations=[item for result in page_results for item in result["violations"]]
    return {"schema_version":BOUNDARY_EVIDENCE_SCHEMA_VERSION,"execution_identity":identity_value,
            "policy_state":{"requested":"enforced","applied":"enforced","effective":"enforced"},
            "source_artifact":_handle(pptx),"rendered_artifact":_handle(pdf),"page_indexes":selected,
            "status":"fail" if violations else "pass","reason_code":BOUNDARY_REASON_CODE if violations else None,
            "violation_count":len(violations),"duplicate_candidate_count":sum(result["duplicate_candidate_count"] for result in page_results),
            "violations":violations}


def build_ir(*, policy: dict[str, Any], frames: list[dict[str, Any]], boundary_analysis: dict[str,Any] | None=None) -> dict[str, Any]:
    declaration=canonicalize_policy(policy)
    by_id={item["component_id"]:item for item in declaration["components"]}
    seen=set(); normalized=[]
    identity_value=None
    for frame in frames:
        validate_frame(frame,declaration)
        component_id=frame["requested"]["component_id"]
        if component_id in seen or component_id not in by_id: raise ValueError("semantic component frame coverage/identity drift")
        seen.add(component_id)
        if identity_value is None: identity_value=frame["execution_identity"]
        elif identity_value!=frame["execution_identity"]: raise ValueError("semantic component execution identity drift")
        normalized.append({"component_id":component_id,"page_index":frame["effective"]["page_index"],
                           "family":frame["effective"]["family"],"topology":frame["effective"]["topology"],
                           "semantic_group_id":frame["effective"]["semantic_group_id"],
                           "member_ids":[item["member_id"] for item in frame["effective"]["members"]],
                           "format":frame["effective"]["format"],"geometry_digest":digest(frame["effective"]["geometry"])})
    if seen!=set(by_id): raise ValueError("semantic component frame coverage incomplete")
    result={"schema_version":IR_SCHEMA_VERSION,"execution_identity":identity_value,"component_count":len(normalized),
            "relation_count":len(declaration["relations"]),"components":sorted(normalized,key=lambda item:item["component_id"]),
            "relations":deepcopy(declaration["relations"]),"status":"pass"}
    if "boundary_no_contact_policy" in declaration:
        if (not isinstance(boundary_analysis,dict) or boundary_analysis.get("schema_version")!=BOUNDARY_EVIDENCE_SCHEMA_VERSION
                or boundary_analysis.get("execution_identity")!=identity_value or boundary_analysis.get("status")!="pass"
                or boundary_analysis.get("violation_count")!=0):
            raise ValueError(BOUNDARY_REASON_CODE)
        result["boundary_no_contact"]={"schema_version":BOUNDARY_EVIDENCE_SCHEMA_VERSION,"status":"pass",
            "source_artifact":deepcopy(boundary_analysis.get("source_artifact")),"rendered_artifact":deepcopy(boundary_analysis.get("rendered_artifact")),
            "page_indexes":deepcopy(boundary_analysis.get("page_indexes",[boundary_analysis.get("page_index")])),
            "violation_count":0,"duplicate_candidate_count":boundary_analysis.get("duplicate_candidate_count",0)}
    elif boundary_analysis is not None:
        raise ValueError("semantic component boundary analysis supplied without declared policy")
    return result


def write_receipt(*, path: Path, policy: dict[str, Any], frames: list[dict[str, Any]], artifacts: list[Path],
                  boundary_analysis: dict[str,Any] | None=None) -> Path:
    ir=build_ir(policy=policy,frames=frames,boundary_analysis=boundary_analysis)
    handles=[{"path":str(item.resolve(strict=True)),"sha256":sha256(item),"size_bytes":item.stat().st_size} for item in artifacts]
    receipt={"schema_version":RECEIPT_SCHEMA_VERSION,"status":"pass","execution_identity":ir["execution_identity"],
             "policy_digest":digest(canonicalize_policy(policy)),"frame_count":len(frames),"ir":ir,"artifacts":handles}
    target=path.resolve(); target.parent.mkdir(parents=True,exist_ok=True); target.write_bytes(_canonical(receipt)+b"\n")
    return target
