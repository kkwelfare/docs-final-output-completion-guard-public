from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from quality_rules import common


def _load_checker():
    spec = importlib.util.spec_from_file_location("composition_receipt_checker", ROOT / "check_docs_artifact_quality.py")
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec); sys.modules[spec.name] = module; spec.loader.exec_module(module)
    return module


def _sha(path: Path) -> str: return hashlib.sha256(path.read_bytes()).hexdigest()
def _handle(path: Path): return {"path": str(path.resolve()), "sha256": _sha(path)}


def _source_and_pdf(tmp_path: Path, *, overlap=False, semantic=False):
    tmp_path.mkdir(parents=True, exist_ok=True)
    pptx = tmp_path / "source.pptx"; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    a = slide.shapes.add_textbox(Pt(30), Pt(30), Pt(150), Pt(24)); a.text_frame.text = "A"
    b = slide.shapes.add_textbox(Pt(40 if overlap else 30), Pt(40 if overlap else 80), Pt(150), Pt(24)); b.text_frame.text = "B"
    if semantic:
        a.name, b.name = "qa-semantic:pair:a", "qa-semantic:pair:b"
    deck.save(pptx)
    pdf = tmp_path / "artifact.pdf"; doc = fitz.open(); page = doc.new_page(width=300, height=180); page.insert_text((30, 40), "English only"); doc.save(pdf); doc.close()
    return pptx, pdf


def _contract(tmp_path: Path, source: Path, artifact: Path):
    suite = Path(__file__)
    rules = {
      "AQ-CHANGE-01": {"approved_scope_ids": ["scope"], "preserved_field_ids": ["preserved"]},
      "AQ-IDENTITY-01": {"identities": []}, "AQ-ASSET-01": {"assets": []},
      "AQ-AUTHORITY-01": {"owner_id": "owner"},
      "AQ-BOUNDARY-01": {"clipped": False, "overflow": False}, "AQ-RELATION-01": {"relations": []},
      "AQ-SCOPE-01": {"protected_region_ids": ["protected"], "baseline_sha256": "a" * 64},
      "AQ-VISIBLE-01": {"criterion_ids": ["v1", "v2", "v3", "v4", "v5"], "image_evidence_sha256": "b" * 64},
      "AQ-FIXTURE-01": {"suite": _handle(suite), "expected_rule_failures": ["none"]},
      "AQ-LOCAL-GLOBAL-01": {"required": False},
    }
    return {"schema_version": "docs-artifact-quality-contract-1", "artifact_kind": "slide_deck", "artifacts": [str(artifact)],
      "authoritative_source": {**_handle(source), "readback_mode": "pptx-geometry"}, "visible_criteria": [{"id": f"v{i}"} for i in range(1, 6)], "rule_contracts": rules,
      "artifact_quality": {"required": True, "layout_typography": {"required": True}, "quality_foundation": {"required": True, "spacing_tokens_pt": {"S": 8}, "semantic_regions": {"minimum_gutter_pt": 0}}, "artifact_profile": {"id": "fixture", "values": {"outer_margin_pt": 0}}}}


def _check(rule: str, contract, artifact):
    spec = contract["rule_contracts"][rule]
    checks = {"AQ-CHANGE-01": {"id": "change-contract", "observed_scope_ids": spec.get("approved_scope_ids"), "changed_preserved_field_ids": []},
      "AQ-IDENTITY-01": {"id": "identity-contract", "identities": []}, "AQ-ASSET-01": {"id": "asset-contract", "assets": []},
      "AQ-AUTHORITY-01": {"id": "authority-contract", "observed_owner_ids": ["owner"], "duplicate_owner_ids": []},
      "AQ-BOUNDARY-01": {"id": "boundary-contract", "clipped": False, "overflow": False, "clipped_count": 0, "overflow_count": 0, "image_cut_count": 0},
      "AQ-RELATION-01": {"id": "relation-contract", "relations": []}, "AQ-SCOPE-01": {"id": "scope-contract", "protected_region_ids": ["protected"], "baseline_sha256": "a" * 64},
      "AQ-VISIBLE-01": {"id": "visible-contract", "criterion_ids": ["v1", "v2", "v3", "v4", "v5"], "image_evidence_sha256": "b" * 64},
      "AQ-FIXTURE-01": {"id": "fixture-contract", "suite": spec.get("suite"), "observed_rule_failures": ["none"]}, "AQ-LOCAL-GLOBAL-01": {"id": "local-global-contract", "required": False}}[rule]
    checks["status"] = "pass"
    return {"schema_version": "docs-artifact-quality-evidence-1", "provenance": {"kind": "checker", "id": "fixture"}, "rule_id": rule, "status": "pass", "review_status": "pass", "artifact_path": str(artifact.resolve()), "artifact_sha256": _sha(artifact), "checks": [checks]}


def test_clean_receipt_passes_l0_to_l3_and_actual_overlap_blocks_l1(tmp_path):
    checker = _load_checker(); source, artifact = _source_and_pdf(tmp_path)
    contract = _contract(tmp_path, source, artifact); contract_path = tmp_path / "contract.json"; contract_path.write_text(json.dumps(contract))
    evidence = tmp_path / "evidence"; evidence.mkdir()
    for rule in common.CONTRACT_RULES:
        (evidence / f"{artifact.name}.{rule}.json").write_text(json.dumps(_check(rule, contract, artifact)))
    receipt, code = checker.build_receipt(contract_path, evidence_dir=evidence)
    validation = {"artifact_path": receipt["entries"][0]["artifact"]["path"], "artifact_sha256": receipt["entries"][0]["artifact"]["sha256"], "actual_delivery_readback": receipt["entries"][0]["actual_delivery_readback"], "rules": receipt["entries"][0]["rules"]}
    optimistic = json.loads(json.dumps(validation)); [rule.update({"status": "pass"}) for rule in optimistic["rules"]]
    assert code == 0 and receipt["status"] == "pass", {"contract": common.validate_contract(contract), "pre_status_entry": common.validate_entry(contract, optimistic)}
    assert {row["layer"] for row in receipt["quality_layers"]} == {"L0", "L1", "L2", "L3"}
    assert all(row["status"] != "block" for row in receipt["quality_layers"])
    composition_path = next(evidence.glob("*.AQ-COMPOSITION-01.json")); assert _sha(composition_path)

    broken_source, broken_artifact = _source_and_pdf(tmp_path / "broken", overlap=True, semantic=True)
    broken = _contract(tmp_path, broken_source, broken_artifact); broken_path = tmp_path / "broken-contract.json"; broken_path.write_text(json.dumps(broken))
    broken_evidence = tmp_path / "broken-evidence"; broken_evidence.mkdir()
    for rule in common.CONTRACT_RULES:
        (broken_evidence / f"{broken_artifact.name}.{rule}.json").write_text(json.dumps(_check(rule, broken, broken_artifact)))
    blocked, exit_code = checker.build_receipt(broken_path, evidence_dir=broken_evidence)
    assert exit_code == 2 and blocked["status"] == "block"
    assert next(row for row in blocked["quality_layers"] if row["rule_id"] == "L1-REGION-01")["status"] == "block"


def test_composition_review_is_coupled_to_hash_bound_layout_receipt(tmp_path):
    checker = _load_checker(); source, artifact = _source_and_pdf(tmp_path, overlap=True)
    contract = _contract(tmp_path, source, artifact); contract_path = tmp_path / "contract.json"; contract_path.write_text(json.dumps(contract))
    evidence = tmp_path / "evidence"; evidence.mkdir()
    for rule in common.CONTRACT_RULES:
        (evidence / f"{artifact.name}.{rule}.json").write_text(json.dumps(_check(rule, contract, artifact)))
    blocked, code = checker.build_receipt(contract_path, evidence_dir=evidence)
    assert code == 2 and blocked["status"] == "block"
    layout = next(rule for rule in blocked["entries"][0]["rules"] if rule["id"] == "AQ-LAYOUT-01")
    scan_path = Path(layout["evidence"]["path"]); scan = json.loads(scan_path.read_text())
    composition_candidates = [item for item in scan["issues"] if item["id"].startswith("composition-")]
    candidate = composition_candidates[0]
    review_path = evidence / f"{artifact.name}.AQ-LAYOUT-01-REVIEW.json"
    review = {"schema_version": "docs-layout-visual-review-1", "status": "pass", "artifact_path": str(artifact.resolve()),
              "artifact_sha256": _sha(artifact), "scan_path": str(scan_path.resolve()), "scan_sha256": scan["reviewed_scan_sha256"],
              "reviews": [{"issue_id": item["id"], "rule": item["rule"], "status": "accepted",
                           "disposition": "accepted", "rationale_code": "verified",
                           "llm_invoked": False, "vision_invoked": False, "page": item["page"],
                           "bbox_pdf_points": item["bbox_pdf_points"], "render_path": item["render_evidence"]["path"], "render_sha256": item["render_evidence"]["sha256"],
                           "affected_pages": item["affected_pages"]} for item in composition_candidates]}
    review_path.write_text(json.dumps(review))
    passed, code = checker.build_receipt(contract_path, evidence_dir=evidence)
    assert code == 0 and passed["status"] == "pass"
    review["reviews"][0]["status"] = "rejected"; review_path.write_text(json.dumps(review))
    failed, code = checker.build_receipt(contract_path, evidence_dir=evidence)
    assert code == 2 and failed["status"] == "block"
